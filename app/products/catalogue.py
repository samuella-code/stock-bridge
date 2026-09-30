"""Shared product validation and bounded spreadsheet ingestion."""
import csv
import io
import zipfile
from decimal import Decimal, InvalidOperation
from sqlalchemy import or_, update
from app import db
from app.models import AuditLog, Business, Product, StockMovement

MAX_ROWS = 1000
MAX_UPLOAD = 2 * 1024 * 1024
HEADERS = {'Product Name': 'name', 'Category': 'category', 'Unit': 'unit',
           'Opening Quantity': 'stock_quantity', 'Cost Price': 'buying_price',
           'Selling Price': 'selling_price', 'Low Stock Threshold': 'minimum_stock_level',
           'SKU': 'sku', 'Description': 'description'}

def product_data(form, editing=False):
    errors, data = [], {}
    for field, limit in [('name',140),('sku',60),('category',80),
                         ('supplier_name',140),('unit',30),('description',500)]:
        value = str(form.get(field) or '').strip()
        if len(value) > limit or any(ord(c) < 32 for c in value):
            errors.append(f"{field.replace('_',' ').title()} is too long or contains control characters.")
        data[field] = value or None
    data['sku'] = data['sku'].upper() if data['sku'] else None
    data['unit'] = data['unit'] or 'unit'
    if not data['name']:
        errors.append('Product name is required.')
    for field in ('buying_price', 'selling_price'):
        try:
            value = Decimal(str(form.get(field) or '0'))
            if not value.is_finite() or value < 0 or value > Decimal('9999999999.99') or value.as_tuple().exponent < -2:
                raise ValueError
            data[field] = value
        except (InvalidOperation, ValueError):
            errors.append(f"{field.replace('_',' ').title()} must be a valid amount with at most two decimal places.")
    fields = [('minimum_stock_level',0),('supplier_lead_time',2),('safety_stock',0)]
    if not editing:
        fields.append(('stock_quantity',0))
    for field, default in fields:
        try:
            value = int(str(form.get(field) or default))
            if value < 0 or value > 2147483647:
                raise ValueError
            data[field] = value
        except (ValueError, TypeError):
            errors.append(f"{field.replace('_',' ').title()} must be a whole number from 0 to 2147483647.")
    return data, errors

def duplicate_errors(bid, data, excluding=None):
    errors = []
    for field in ('sku',):
        value = data.get(field)
        if value:
            query = Product.query.filter(Product.business_id == bid, getattr(Product, field) == value)
            if excluding is not None:
                query = query.filter(Product.id != excluding)
            if query.first():
                errors.append('That SKU is already used in this business (including archived products).')
    return errors

def validate_rows(bid, rows):
    if not rows or len(rows) > MAX_ROWS:
        raise ValueError(f'Enter between 1 and {MAX_ROWS} products per batch.')
    normalized, issues = [], []
    # One SKU query, rather than one query per uploaded product.
    skus = {str(r.get('sku') or '').strip().upper() for r in rows} - {''}
    existing_skus = {r[0] for r in db.session.query(Product.sku).filter(Product.business_id == bid, Product.sku.in_(skus))} if skus else set()
    seen_skus, seen_rows = set(), set()
    for number, raw in enumerate(rows, 2):
        data, errors = product_data(raw)
        for field, seen, existing in [('sku',seen_skus,existing_skus)]:
            value = data.get(field)
            if value:
                if value in seen or value in existing:
                    errors.append('Duplicate SKU.')
                seen.add(value)
        key = tuple(sorted((k,str(v).casefold() if k == 'name' else str(v)) for k,v in data.items()))
        if key in seen_rows:
            errors.append('Duplicate product row.')
        seen_rows.add(key)
        normalized.append(data)
        if errors:
            issues.append({'row':raw.get('_row', number), 'errors':errors})
    return normalized, issues

def add_product(bid, data):
    data = dict(data)
    qty = data.pop('stock_quantity')
    product = Product(business_id=bid, stock_quantity=qty, opening_quantity=qty, **data)
    db.session.add(product)
    db.session.flush()
    db.session.add(StockMovement(business_id=bid, product_id=product.id, kind='opening',
                                quantity_change=qty, occurred_at=product.created_at))
    return product

def save_batch(bid, rows, actor_id, receipt=None):
    # Serialise confirmations for a business. A no-op UPDATE obtains a write lock
    # on both SQLite and Postgres; everything, including the receipt, commits once.
    db.session.execute(update(Business).where(Business.id == bid).values(name=Business.name))
    if receipt and AuditLog.query.filter_by(business_id=bid, action='PRODUCT_IMPORT', description=receipt).first():
        db.session.rollback()
        return 0
    normalized, issues = validate_rows(bid, rows) if rows else ([], [])
    if not rows and not receipt:
        db.session.rollback()
        raise ValueError('Enter at least one product.')
    if issues:
        db.session.rollback()
        raise ValueError('Some rows changed or have duplicate identifiers. Preview the products again.')
    try:
        products = []
        for data in normalized:
            data = dict(data)
            quantity = data.pop('stock_quantity')
            products.append(Product(business_id=bid, stock_quantity=quantity, opening_quantity=quantity, **data))
        db.session.add_all(products)
        db.session.flush()
        db.session.add_all([StockMovement(business_id=bid, product_id=p.id, kind='opening',
            quantity_change=p.opening_quantity, occurred_at=p.created_at) for p in products])
        db.session.add(AuditLog(actor_id=actor_id, business_id=bid, action='PRODUCT_IMPORT',
                               description=receipt or f'Quick Add: {len(rows)} products'))
        db.session.commit()
    except Exception:
        db.session.rollback()
        raise
    return len(rows)

def search_products(bid, term):
    query = Product.query.filter_by(business_id=bid, active=True)
    if term:
        # Escape wildcard characters so identifiers are literal, not patterns.
        escaped = term.replace('\\','\\\\').replace('%','\\%').replace('_','\\_')
        pattern = f'%{escaped}%'
        query = query.filter(or_(Product.name.ilike(pattern, escape='\\'), Product.sku.ilike(pattern, escape='\\'),
                                 Product.category.ilike(pattern, escape='\\')))
    return query

def read_spreadsheet(upload):
    filename = (upload.filename or '').lower()
    blob = upload.stream.read(MAX_UPLOAD + 1)
    if len(blob) > MAX_UPLOAD:
        raise ValueError('The spreadsheet must be smaller than 2 MB. Split it into smaller files.')
    if filename.endswith('.csv'):
        try:
            csv.field_size_limit(10000)
            grid = csv.reader(io.StringIO(blob.decode('utf-8-sig')), strict=True)
            return grid_to_rows(grid)
        except (UnicodeError, csv.Error):
            raise ValueError('This CSV could not be read. Save it as UTF-8 CSV and try again.')
    if not filename.endswith('.xlsx'):
        raise ValueError('Upload a CSV or .xlsx spreadsheet. Old .xls files are not supported.')
    try:
        with zipfile.ZipFile(io.BytesIO(blob)) as archive:
            entries = archive.infolist()
            if len(entries) > 200 or sum(e.file_size for e in entries) > 10 * 1024 * 1024:
                raise ValueError('The expanded spreadsheet is too large. Use the StockBridge template.')
            if any('vbaProject' in e.filename or 'externalLinks/' in e.filename for e in entries):
                raise ValueError('Macros and external workbook links are not allowed.')
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(blob), read_only=True, data_only=False, keep_links=False)
        try:
            if len(workbook.worksheets) != 1:
                raise ValueError('Use one worksheet per import.')
            sheet = workbook.worksheets[0]
            if sheet.max_column and sheet.max_column > len(HEADERS) + 1:
                raise ValueError('Use only the columns in the StockBridge template.')
            sheet.reset_dimensions()
            def values():
                for cells in sheet.iter_rows():
                    if any(c.data_type == 'f' for c in cells):
                        raise ValueError('Spreadsheet formulas are not allowed. Paste values instead.')
                    yield [c.value for c in cells]
            return grid_to_rows(values())
        finally:
            workbook.close()
    except ValueError:
        raise
    except Exception:
        raise ValueError('This spreadsheet is malformed or cannot be read. Use the StockBridge template.')

def grid_to_rows(grid):
    iterator = iter(grid)
    headers = next(iterator, [])
    headers = [str(h or '').strip().removesuffix(' *') for h in headers]
    lookup = {k.lower():v for k,v in HEADERS.items()}
    # Backward compatibility: discard this optional column from older templates.
    # Never validate, save or overwrite retained legacy identifier values.
    lookup['barcode'] = None
    if not headers or len(set(h.lower() for h in headers)) != len(headers) or any(h.lower() not in lookup for h in headers):
        raise ValueError('Use the StockBridge template column names; duplicate or unknown headers are not allowed.')
    fields = [lookup[h.lower()] for h in headers]
    if 'name' not in fields:
        raise ValueError('The Product Name column is required.')
    rows = []
    physical_rows = 0
    for cells in iterator:
        physical_rows += 1
        if physical_rows > MAX_ROWS:
            raise ValueError(f'Use at most {MAX_ROWS} spreadsheet rows per import; split larger catalogues into batches.')
        if not any(v is not None and str(v).strip() for v in cells):
            continue
        if len(cells) > len(fields) and any(str(v or '').strip() for v in cells[len(fields):]):
            raise ValueError('A row has extra columns. Check the template and CSV separators.')
        raw = {'_row': physical_rows + 1}
        for field, value in zip(fields, cells):
            if field is None:
                continue
            if field == 'sku' and value is not None and not isinstance(value,str):
                raise ValueError('Format the SKU column as Text to preserve leading zeros.')
            text = '' if value is None else str(value)
            if text.lstrip().startswith(('=','+','@')):
                raise ValueError('Formula-like cells are not allowed. Paste plain values instead.')
            raw[field] = text
        rows.append(raw)
    if not rows:
        raise ValueError('No products found. Fill in the template before uploading.')
    return rows
